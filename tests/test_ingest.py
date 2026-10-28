"""M5 解析、清洗与去重的单元测试。不联网、不依赖 .cache。"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from medidiag.rag.ingest.clean import clean_document, normalize_text
from medidiag.rag.ingest.dedup import dedup_documents, dedup_paragraphs
from medidiag.rag.ingest.document import (
    Document,
    DocumentMetadata,
    Element,
    Section,
    serialize_table_row,
)
from medidiag.rag.ingest.encoding import decode_bytes
from medidiag.rag.ingest.html import extract_who_page_date, parse_html, parse_mediawiki_json
from medidiag.rag.ingest.pdf import (
    classify_heading,
    detect_header_footer,
    merge_broken_lines,
    remove_header_footer,
)

# ---------- document.py ----------


def test_document_json_roundtrip() -> None:
    doc = Document(
        doc_id="d1",
        metadata=DocumentMetadata(source_type="who_factsheet", title="季节性流感"),
        sections=[
            Section(title="传播", level=2, elements=[
                Element(kind="paragraph", text="流感通过飞沫传播。"),
                Element(kind="list", text="a\nb", items=["戴口罩", "勤洗手"]),
                Element(kind="table", text="列1：值1", header=["列1"], rows=[["值1"]]),
            ]),
        ],
    )
    loaded = Document.from_json(doc.to_json())
    assert loaded.doc_id == doc.doc_id
    assert loaded.metadata.source_type == "who_factsheet"
    assert loaded.sections[0].elements[1].items == ["戴口罩", "勤洗手"]
    assert loaded.sections[0].elements[2].header == ["列1"]


def test_section_paths() -> None:
    doc = Document(doc_id="d", metadata=DocumentMetadata(), sections=[
        Section(title="一、总则", level=1),
        Section(title="（一）定义", level=2),
        Section(title="（二）范围", level=2),
        Section(title="二、诊疗", level=1),
    ])
    paths = [p for _, p in doc.iter_section_paths()]
    assert paths[1] == ["一、总则", "（一）定义"]
    assert paths[3] == ["二、诊疗"]


def test_serialize_table_row() -> None:
    assert serialize_table_row(["地区", "病例数"], ["广东", "120"]) == "地区：广东；病例数：120"
    # 缺表头时按列N命名；空单元格跳过
    assert serialize_table_row([], ["广东", ""]) == "列1：广东"


# ---------- encoding.py ----------


def test_decode_gbk_bytes() -> None:
    raw = "流行性感冒诊疗方案".encode("gb18030")
    result = decode_bytes(raw)
    assert result.text == "流行性感冒诊疗方案"
    assert result.encoding == "gb18030"


def test_decode_respects_meta_declared() -> None:
    body = "<html><head><meta charset='gbk'></head><body>感冒</body></html>".encode("gb18030")
    result = decode_bytes(body)
    assert "感冒" in result.text
    assert result.source == "declared"


def test_decode_utf8_fallback() -> None:
    result = decode_bytes("冠状病毒病".encode())
    assert result.text == "冠状病毒病"


# ---------- html.py ----------


def test_html_heading_hierarchy_and_table() -> None:
    html = """
    <html><body>
    <article>
    <h1>季节性流感</h1>
    <p>流感是急性呼吸道感染。</p>
    <h2>传播</h2>
    <p>通过飞沫传播。</p>
    <h3>高危人群</h3>
    <p>老年人风险更高。</p>
    <h2>预防</h2>
    <table><thead><tr><th>措施</th><th>说明</th></tr></thead>
    <tbody><tr><td>疫苗</td><td>每年接种</td></tr></tbody></table>
    </article>
    </body></html>
    """
    doc = parse_html(html, "flu", source_type="who_factsheet")
    titles = [(s.title, s.level) for s in doc.sections]
    assert ("季节性流感", 1) in titles
    assert ("传播", 2) in titles
    assert ("高危人群", 3) in titles
    assert ("预防", 2) in titles
    tables = [e for s in doc.sections for e in s.elements if e.kind == "table"]
    assert len(tables) == 1
    assert tables[0].header == ["措施", "说明"]
    assert tables[0].rows == [["疫苗", "每年接种"]]
    assert tables[0].text == "措施：疫苗；说明：每年接种"
    # 章节路径
    paths = {s.title: p for s, p in doc.iter_section_paths() if s.title}
    assert paths["高危人群"] == ["季节性流感", "传播", "高危人群"]


def test_html_skips_nav_script_share() -> None:
    html = """
    <html><body>
    <nav>面包屑导航</nav>
    <script>var x = 1;</script>
    <div class="social-share">分享控件文字</div>
    <article><h1>标题</h1><p>正文内容，应当保留。</p></article>
    <footer>页脚</footer>
    </body></html>
    """
    doc = parse_html(html, "d", source_type="who_factsheet")
    text = doc.full_text()
    assert "正文内容" in text
    assert "面包屑" not in text
    assert "var x" not in text
    assert "分享控件" not in text
    assert "页脚" not in text


def test_html_list_items() -> None:
    html = "<article><h2>重要事实</h2><ul><li>第一条事实。</li><li>第二条事实。</li></ul></article>"
    doc = parse_html(html, "d", source_type="who_factsheet")
    lists = [e for s in doc.sections for e in s.elements if e.kind == "list"]
    assert len(lists) == 1
    assert lists[0].items == ["第一条事实。", "第二条事实。"]


def test_mediawiki_json_parse() -> None:
    payload = {
        "parse": {
            "title": "流行性感冒",
            "text": {"*": "<h2>症状</h2><p>发热、咳嗽。<span class=\"mw-editsection\">[编辑]</span></p>"
                       "<div class=\"navbox\">导航框噪声</div><table><tr><th>项目</th></tr>"
                       "<tr><td>数据</td></tr></table>"},
        }
    }
    doc = parse_mediawiki_json(payload, "wiki_flu")
    text = doc.full_text()
    assert doc.metadata.title == "流行性感冒"
    assert "发热、咳嗽" in text
    assert "导航框噪声" not in text
    assert "[编辑]" not in text
    tables = [e for s in doc.sections for e in s.elements if e.kind == "table"]
    assert len(tables) == 1


# ---------- pdf.py 纯函数 ----------


def test_detect_header_footer() -> None:
    pages = [
        ["世界卫生组织", "正文第一段", "第 1 页"],
        ["世界卫生组织", "正文第二段", "第 2 页"],
        ["世界卫生组织", "正文第三段", "第 3 页"],
    ]
    repeated = detect_header_footer(pages, min_pages=3)
    assert "世界卫生组织" in repeated
    assert "第 1 页" in repeated
    cleaned = remove_header_footer(pages, repeated)
    assert cleaned[0] == ["正文第一段"]


def test_detect_header_footer_short_doc() -> None:
    # 不足 min_pages 页时不做跨页判定，避免误删
    pages = [["标题", "正文"], ["标题", "正文二"]]
    assert detect_header_footer(pages, min_pages=3) == {"标题"} - {"标题"} or True
    repeated = detect_header_footer(pages, min_pages=3)
    assert "标题" not in repeated


def test_merge_broken_lines() -> None:
    lines = [
        "流行性感冒是由流感病毒引起的急性呼吸道",
        "传染病，主要通过飞沫传播。",
        "一、病原学",
        "流感病毒属于正黏病毒科，",
        "分为甲、乙、丙三型。",
    ]
    merged = merge_broken_lines(lines)
    assert merged[0] == "流行性感冒是由流感病毒引起的急性呼吸道传染病，主要通过飞沫传播。"
    assert merged[1] == "一、病原学"
    assert merged[2] == "流感病毒属于正黏病毒科，分为甲、乙、丙三型。"


def test_merge_broken_lines_english_space() -> None:
    merged = merge_broken_lines(["COVID-19 is caused by", "SARS-CoV-2 virus."])
    assert merged == ["COVID-19 is caused by SARS-CoV-2 virus."]


def test_classify_heading_by_numbering() -> None:
    assert classify_heading("一、病原学") == 1
    assert classify_heading("（一）流行病学") == 2
    assert classify_heading("(二) 临床表现") == 2
    assert classify_heading("1. 诊断标准") == 3
    assert classify_heading("（3）鉴别诊断") == 4
    assert classify_heading("这是一条普通正文，没有编号，长度足以超过标题限制很多很多的字。" * 2) == 0
    assert classify_heading("普通正文没有编号。") == 0


def test_classify_heading_by_font_size() -> None:
    assert classify_heading("临床表现", font_size=16.0, body_size=12.0) == 2
    assert classify_heading("临床表现", font_size=12.0, body_size=12.0) == 0


def test_classify_heading_tiao_vs_zhang() -> None:
    # 「第X章/节/篇/部分」是文档结构标题；「第X条」是条款正文开头，不作标题。
    # 两份真实 DOCX（门诊信息页、乡镇卫生院标准）无「第X条」，规则按通用公文约定。
    assert classify_heading("第三章 预防") == 1
    assert classify_heading("第3节 诊断") == 1
    assert classify_heading("第一篇 总则") == 1
    assert classify_heading("第三条 医疗机构应当按照规定填写诊疗信息") == 0
    assert classify_heading("第十条 医师应当签字") == 0
    assert classify_heading("第3条 本规定自发布之日起施行") == 0


# ---------- docx / xlsx / csv（程序化生成小文件） ----------


def test_parse_docx(tmp_path: Path) -> None:
    import docx

    from medidiag.rag.ingest.docx import parse_docx

    doc = docx.Document()
    doc.add_heading("流行性感冒诊疗方案", level=1)
    doc.add_paragraph("本方案适用于各级医疗机构。")
    doc.add_heading("一、病原学", level=2)
    doc.add_paragraph("流感病毒属于正黏病毒科。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "型别"
    table.cell(0, 1).text = "特点"
    table.cell(1, 0).text = "甲型"
    table.cell(1, 1).text = "易变异"
    path = tmp_path / "flu.docx"
    doc.save(str(path))

    parsed, report = parse_docx(path, "flu_scheme")
    titles = [(s.title, s.level) for s in parsed.sections]
    assert ("流行性感冒诊疗方案", 1) in titles
    assert ("一、病原学", 2) in titles
    assert report.table_count == 1
    tables = [e for s in parsed.sections for e in s.elements if e.kind == "table"]
    assert tables[0].header == ["型别", "特点"]
    assert tables[0].text == "型别：甲型；特点：易变异"


def test_parse_docx_numbered_heading_without_style(tmp_path: Path) -> None:
    import docx

    from medidiag.rag.ingest.docx import parse_docx

    doc = docx.Document()
    doc.add_paragraph("一、总则")
    doc.add_paragraph("正文内容。")
    path = tmp_path / "plain.docx"
    doc.save(str(path))
    parsed, _ = parse_docx(path, "d")
    assert any(s.title == "一、总则" for s in parsed.sections)


def test_parse_docx_township_pattern(tmp_path: Path) -> None:
    """按真实文件 nhc_township_hospital_standard_attach1 的结构生成：

    无 Heading 样式，靠「一、」「二、」编号组织病种清单，表格按科室多级组织。
    """
    import docx

    from medidiag.rag.ingest.docx import parse_docx

    doc = docx.Document()
    doc.add_paragraph("附件1")
    doc.add_paragraph("门诊及收治住院基本病种")
    doc.add_paragraph("一、基本病种（128种）")
    t1 = doc.add_table(rows=3, cols=3)
    for j, h in enumerate(["序号", "病种", "ICD代码"]):
        t1.cell(0, j).text = h
    t1.cell(1, 0).text = "一、内科（42种）"
    t1.cell(2, 1).text = "高血压病"
    t1.cell(2, 2).text = "I10.x00"
    doc.add_paragraph("二、中医疾病名称（66种）")
    t2 = doc.add_table(rows=2, cols=3)
    for j, h in enumerate(["序号", "病种", "TCD代码"]):
        t2.cell(0, j).text = h
    path = tmp_path / "township.docx"
    doc.save(str(path))

    parsed, report = parse_docx(path, "township")
    titles = [(s.title, s.level) for s in parsed.sections]
    # 「一、基本病种」「二、中医疾病名称」识别为 1 级标题
    assert ("一、基本病种（128种）", 1) in titles
    assert ("二、中医疾病名称（66种）", 1) in titles
    assert report.table_count == 2
    tables = [e for s in parsed.sections for e in s.elements if e.kind == "table"]
    assert tables[0].header == ["序号", "病种", "ICD代码"]


def test_parse_xlsx(tmp_path: Path) -> None:
    import openpyxl

    from medidiag.rag.ingest.table import parse_xlsx

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026年9月"
    ws.append(["地区", "发病数", "死亡数"])
    ws.append(["广东", 1200, 3])
    ws.append(["北京", 800, 1])
    path = tmp_path / "epi.xlsx"
    wb.save(str(path))

    parsed, report = parse_xlsx(path, "epi_monthly")
    assert parsed.sections[0].title == "2026年9月"
    rows = [e for s in parsed.sections for e in s.elements
            if e.kind == "paragraph" and e.metadata.get("content_kind") == "table_row"]
    assert rows[0].text == "地区：广东；发病数：1200；死亡数：3"
    assert rows[0].metadata["row"] == 2
    summaries = [e for s in parsed.sections for e in s.elements if e.kind == "table"]
    assert len(summaries) == 1
    assert "共 2 行" in summaries[0].text
    assert report.sheet_count == 1


def test_parse_csv(tmp_path: Path) -> None:
    from medidiag.rag.ingest.table import parse_csv

    path = tmp_path / "data.csv"
    path.write_text("地区,发病数\n广东,1200\n北京,800\n", encoding="utf-8")
    parsed, report = parse_csv(path, "csv_doc")
    rows = [e for s in parsed.sections for e in s.elements if e.kind == "paragraph"]
    assert rows[1].text == "地区：北京；发病数：800"
    assert report.row_count == 2


# ---------- 数据质量修复：元数据回填与页首杂质 ----------


def test_who_factsheet_backfills_title_and_published_at() -> None:  # noqa: E501
    """WHO 事实清单：h1 回填 title，页首日期行提取为 published_at。"""
    html = """
    <html><head><title>麻疹</title></head><body>
    <article>
    <div class="sf-item-header-wrapper">
      <h1>麻疹</h1>
      <div class="date"><span class="timestamp">2026年7月15日</span></div>
    </div>
    <p>麻疹是一种传染性很强的病毒性疾病。</p>
    <h2>重要事实</h2>
    <ul><li>麻疹可通过接种疫苗预防。</li></ul>
    </article>
    </body></html>
    """
    doc = parse_html(html, "who_fs_measles", source_type="who_factsheet")
    assert doc.metadata.title == "麻疹"
    assert doc.metadata.published_at == "2026-07-15"
    assert "传染性很强" in doc.full_text()


def test_extract_who_page_date() -> None:
    assert extract_who_page_date('<span class="timestamp">2025年12月8日</span>') == "2025-12-08"
    assert extract_who_page_date("<p>无日期</p>") is None


def test_who_header_noise_and_cr_removed_by_clean() -> None:
    """WHO 页首杂质（跳转链接/日期行/栏目标签）与 \\r 在清洗后不存在。"""
    doc = Document(doc_id="d", metadata=DocumentMetadata(source_type="who_factsheet"), sections=[
        Section(title="", level=1, elements=[
            Element(kind="paragraph", text="麻疹\r\n \r\n Skip to main content"),
            Element(kind="paragraph", text="2026年7月15日"),
            Element(kind="paragraph", text="实况报道"),
            Element(kind="paragraph", text="麻疹是病毒性疾病。"),
        ]),
    ])
    clean_document(doc)
    text = doc.full_text()
    assert "\r" not in text
    assert "Skip to main content" not in text
    assert "麻疹是病毒性疾病" in text


def test_clean_drops_reference_section_by_source() -> None:
    """参考文献/外部链接类章节按来源整节丢弃，正文章节不受影响。"""
    for stype in ("who_factsheet", "wikipedia", "nhc_document"):
        doc = Document(doc_id="d", metadata=DocumentMetadata(source_type=stype), sections=[
            Section(title="概述", level=2, elements=[
                Element(kind="paragraph", text="麻疹是传染性很强的病毒性疾病，可接种疫苗预防。"),
            ]),
            Section(title="参考文献", level=2, elements=[
                Element(kind="paragraph", text="1. 《柳叶刀》2024年；403:2162-2203。"),
            ]),
            Section(title="外部链接", level=2, elements=[
                Element(kind="paragraph", text="https://example.org/more"),
            ]),
        ])
        clean_document(doc)
        titles = [s.title for s in doc.sections]
        assert "概述" in titles
        assert "参考文献" not in titles
        assert "外部链接" not in titles
        assert "传染性很强" in doc.full_text()


def test_clean_drops_page_number_lines_who_publication() -> None:
    """WHO 出版物：罗马数字页码孤立行删除；无标题参考文献列表（连续编号+年份/URL）丢弃。"""
    doc = Document(doc_id="d", metadata=DocumentMetadata(source_type="who_publication"), sections=[
        Section(title="前言", level=1, elements=[
            Element(kind="paragraph", text="vi"),
            Element(kind="paragraph", text="本指南旨在提供建议。"),
        ]),
        Section(title="正文", level=1, elements=[
            Element(kind="paragraph", text="第一条建议内容，足够长的正文段落，不应被误删。"),
            Element(kind="paragraph", text="1. 世界卫生组织；2024年。https://iris.who.int/handle/10665/68571。"),
            Element(kind="paragraph", text="2. 日内瓦：世界卫生组织；2023年。https://iris.who.int/handle/10665/272452。"),
            Element(kind="paragraph", text="3. 《柳叶刀》2024年；403:2162-2203。"),
        ]),
    ])
    clean_document(doc)
    text = doc.full_text()
    assert "vi" not in [e.text for s in doc.sections for e in s.elements]
    assert "本指南旨在提供建议" in text
    assert "第一条建议内容" in text
    # 连续 3 行无标题参考文献被丢弃
    assert "iris.who.int" not in text
    assert "2162-2203" not in text


def test_clean_drops_untitled_english_bibliography_who_factsheet() -> None:
    """WHO 事实清单：正文末尾章节里连续 2 行以上的英文书目按参考文献丢弃。

    覆盖「世卫组织的应对」章节末尾带英文参考文献、无「参考文献」标题的形态。
    """
    doc = Document(doc_id="d", metadata=DocumentMetadata(source_type="who_factsheet"), sections=[
        Section(title="世卫组织的应对", level=2, elements=[
            Element(kind="paragraph", text="世卫组织提供全球技术指导，正文段落应保留。"),
            Element(kind="paragraph", text="(1) Bearak J, Popinchalk A, et al. Lancet Glob Health. 2020 Sep."),
            Element(kind="paragraph", text="(2) Ganatra B, et al. doi: 10.1016/S2214-109X(20)30315-6."),
        ]),
    ])
    clean_document(doc)
    text = doc.full_text()
    assert "世卫组织提供全球技术指导" in text
    assert "Bearak" not in text
    assert "Ganatra" not in text


def test_clean_keeps_short_untitled_reference_runs() -> None:
    """连续不足 3 行的编号引用不判为参考文献（避免误删正文编号段落）。"""
    doc = Document(doc_id="d", metadata=DocumentMetadata(source_type="who_publication"), sections=[
        Section(title="正文", level=1, elements=[
            Element(kind="paragraph", text="1. 世界卫生组织2024年报告提到相关建议。"),
            Element(kind="paragraph", text="后续正文段落，内容足够长，不应被误删。"),
        ]),
    ])
    clean_document(doc)
    assert "世界卫生组织2024年报告" in doc.full_text()


def test_epidemic_report_title_backfilled_from_h5() -> None:
    """疫情月报：标题（<h5>）在正文容器外，按来源规则回填 metadata.title。"""
    html = """
    <html><body>
    <div class="containerBox">
      <h5><a>2024年10月全国法定传染病疫情概况</a></h5>
      <div class="xqCon"><span class="fb">时间：<em>2024-11-20</em></span></div>
      <div class="content" id="articleCon"><p>2024年10月，全国共报告法定传染病563572例。</p></div>
    </div>
    </body></html>
    """
    doc = parse_html(html, "cdc_report_202410", source_type="epidemic_report")
    assert doc.metadata.title == "2024年10月全国法定传染病疫情概况"
    assert doc.metadata.published_at == "2024-11-20"
    assert "563572例" in doc.full_text()


def test_pipeline_backfills_accessed_at_from_fetch_report(tmp_path: Path) -> None:
    """pipeline 从 fetch_report.json 回填 accessed_at；清单显式值优先。"""
    from medidiag.rag.ingest.pipeline import run_ingest

    raw = tmp_path / "raw" / "who_factsheet"
    raw.mkdir(parents=True)
    (raw / "flu.html").write_text(
        "<article><h1>季节性流感</h1><p>流感通过飞沫传播。</p></article>", encoding="utf-8")
    sources = {
        "version": "test",
        "source_types": {"who_factsheet": {"publisher": "世界卫生组织"}},
        "documents": [
            {"doc_id": "flu", "source_type": "who_factsheet", "title": "季节性流感",
             "url": "https://example.org/flu", "format": "html", "license_conclusion": "纳入"},
        ],
    }
    (tmp_path / "sources.json").write_text(
        json.dumps(sources, ensure_ascii=False), encoding="utf-8")
    fetch_report = {"records": [
        {"doc_id": "flu", "accessed_at": "2026-10-27T08:39:16Z"},
    ]}
    (tmp_path / "fetch_report.json").write_text(
        json.dumps(fetch_report, ensure_ascii=False), encoding="utf-8")

    report = run_ingest(tmp_path / "sources.json", tmp_path)
    assert report.files[0].status == "ok"
    line = (tmp_path / "parsed" / "documents.jsonl").read_text(encoding="utf-8").splitlines()[0]
    loaded = Document.from_json(line)
    assert loaded.metadata.accessed_at == "2026-10-27T08:39:16Z"


# ---------- clean.py ----------


def test_normalize_text_fullwidth() -> None:
    # 全角字母/数字/空格转半角；中文全角标点（，：（）％等）保持全角
    assert normalize_text("（一）　感冒，发热　50％") == "（一） 感冒，发热 50％"
    assert normalize_text("ＡＢＣ　１２３") == "ABC 123"


def test_clean_removes_reference_markers() -> None:
    doc = Document(doc_id="d", metadata=DocumentMetadata(source_type="who_factsheet"), sections=[
        Section(title="概述", level=2, elements=[
            Element(kind="paragraph", text="全球每年约发生7300万例(1)。人工流产常见(2)。"),
        ]),
    ])
    clean_document(doc)
    text = doc.sections[0].elements[0].text
    assert "(1)" not in text
    assert "(2)" not in text
    assert "7300万例" in text


def test_clean_drops_control_lines_by_source() -> None:
    doc = Document(doc_id="d", metadata=DocumentMetadata(source_type="wikipedia"), sections=[
        Section(title="", level=1, elements=[
            Element(kind="paragraph", text="[编辑]"),
            Element(kind="paragraph", text="流行性感冒是传染病。"),
        ]),
    ])
    clean_document(doc)
    texts = [e.text for s in doc.sections for e in s.elements]
    assert texts == ["流行性感冒是传染病。"]


def test_normalize_text_unifies_cr() -> None:
    # \r\n / \r 统一为 \n；全半角归一不受影响
    assert normalize_text("麻疹\r\n \r\n 概述") == "麻疹\n \n 概述"
    assert normalize_text("a\rb") == "a\nb"


def test_clean_keeps_table_elements() -> None:
    doc = Document(doc_id="d", metadata=DocumentMetadata(source_type="who_factsheet"), sections=[
        Section(title="t", level=1, elements=[
            Element(kind="table", text="地区：广东(1)", header=["地区"], rows=[["广东(1)"]]),
        ]),
    ])
    clean_document(doc)
    assert doc.sections[0].elements[0].text == "地区：广东(1)"


# ---------- dedup.py ----------


def _doc(doc_id: str, text: str, level: str) -> Document:
    return Document(doc_id=doc_id, metadata=DocumentMetadata(evidence_level=level), sections=[
        Section(title="", level=1, elements=[Element(kind="paragraph", text=text)]),
    ])


def test_dedup_keeps_higher_authority() -> None:
    text = "流感是由流感病毒引起的急性呼吸道传染病，主要通过飞沫传播，每年季节性流行。" * 3
    low = _doc("low", text, "level_5_other")
    high = _doc("high", text, "level_1_guideline")
    report = dedup_documents([low, high])
    assert report.kept_doc_ids == ["high"]
    assert report.dropped_doc_ids == ["low"]
    assert report.pairs[0].kept_id == "high"
    assert report.pairs[0].granularity == "document"


def test_dedup_distinct_docs_kept() -> None:
    a = _doc("a", "流感通过飞沫传播，预防靠疫苗接种与勤洗手，流行于冬春季。" * 3, "level_5_other")
    b = _doc("b", "糖尿病是慢性代谢性疾病，特征为长期高血糖，需要综合管理。" * 3, "level_5_other")
    report = dedup_documents([a, b])
    assert set(report.kept_doc_ids) == {"a", "b"}
    assert not report.dropped_doc_ids


def test_dedup_paragraphs_cross_source() -> None:
    shared = "接种疫苗是预防流感最有效的手段，可以显著降低接种者罹患流感和发生严重并发症的风险。"
    doc_a = Document(doc_id="a", metadata=DocumentMetadata(evidence_level="level_1_guideline"),
                     sections=[Section(title="", level=1, elements=[
                         Element(kind="paragraph", text=shared)])])
    doc_b = Document(doc_id="b", metadata=DocumentMetadata(evidence_level="level_5_other"),
                     sections=[Section(title="", level=1, elements=[
                         Element(kind="paragraph", text=shared),
                         Element(kind="paragraph", text="流感与普通感冒不同，病原完全不同。")])])
    report = dedup_paragraphs([doc_b, doc_a])
    assert report.dropped_paragraphs.get("b") == [shared]
    assert "a" not in report.dropped_paragraphs


# ---------- pipeline ----------


def test_run_ingest_end_to_end(tmp_path: Path) -> None:
    from medidiag.rag.ingest.pipeline import run_ingest

    raw = tmp_path / "raw" / "who_factsheet"
    raw.mkdir(parents=True)
    (raw / "flu.html").write_text(
        "<article><h1>季节性流感</h1><p>流感通过飞沫传播。</p></article>", encoding="utf-8")
    sources = {
        "version": "test",
        "source_types": {"who_factsheet": {"publisher": "世界卫生组织", "evidence_level": "level_5_other"}},
        "documents": [
            {"doc_id": "flu", "source_type": "who_factsheet", "title": "季节性流感",
             "url": "https://example.org/flu", "format": "html", "language": "zh-CN",
             "license_conclusion": "纳入", "evidence_level": "level_5_other"},
            {"doc_id": "excluded", "source_type": "who_factsheet", "title": "不纳入",
             "url": "https://example.org/no", "format": "html", "language": "zh-CN",
             "license_conclusion": "排除", "evidence_level": "level_5_other"},
        ],
    }
    sources_path = tmp_path / "sources.json"
    sources_path.write_text(json.dumps(sources, ensure_ascii=False), encoding="utf-8")

    report = run_ingest(sources_path, tmp_path)
    assert len(report.files) == 1
    assert report.files[0].status == "ok"
    assert report.files[0].element_count >= 1
    docs_out = (tmp_path / "parsed" / "documents.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(docs_out) == 1
    loaded = Document.from_json(docs_out[0])
    assert loaded.doc_id == "flu"
    assert loaded.metadata.evidence_level == "level_5_other"
    assert loaded.metadata.content_hash == report.files[0].content_hash
    assert (tmp_path / "parsed" / "report.json").exists()


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))
